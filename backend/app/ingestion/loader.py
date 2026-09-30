import hashlib
import io
from pathlib import Path
from typing import Tuple
import fitz
from pptx import Presentation

from backend.app.config import settings
from backend.app.core.security import sanitize_filename
from backend.app.core.errors import ValidationError

def compute_sha256(content: bytes) -> str:
    """Calculate SHA-256 hash of raw file bytes for deduplication."""
    return hashlib.sha256(content).hexdigest()

def convert_to_pdf(filename: str, content: bytes) -> bytes:
    """Convert various file formats to PDF for unified ingestion."""
    ext = filename.lower().split('.')[-1]
    if ext == "pdf":
        return content
    elif ext in ["png", "jpg", "jpeg", "webp", "tiff"]:
        doc = fitz.open(stream=content, filetype=ext)
        pdfbytes = doc.convert_to_pdf()
        doc.close()
        return pdfbytes
    elif ext == "txt":
        doc = fitz.open(stream=content, filetype="txt")
        pdfbytes = doc.convert_to_pdf()
        doc.close()
        return pdfbytes
    elif ext == "pptx":
        prs = Presentation(io.BytesIO(content))
        doc = fitz.open()
        for slide in prs.slides:
            text = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text.append(shape.text)
            page = doc.new_page()
            # Simple text insertion; could be styled better, but suffices for indexing
            p = fitz.Point(50, 50)
            page.insert_text(p, "\\n".join(text), fontsize=11)
        pdfbytes = doc.tobytes()
        doc.close()
        return pdfbytes
    else:
        # Try PyMuPDF's general conversion as a fallback
        try:
            doc = fitz.open(stream=content, filetype=ext)
            pdfbytes = doc.convert_to_pdf()
            doc.close()
            return pdfbytes
        except Exception:
            raise ValidationError(f"Unsupported file format: {ext}")

def save_uploaded_document(workspace_id: str, filename: str, content: bytes) -> Tuple[Path, str, str, int]:
    """Validate, sanitize, convert to PDF if needed, and save to workspace storage.
    
    Returns:
        (saved_path, sanitized_filename, sha256_hash, file_size_bytes)
    """
    max_bytes = settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024
    if len(content) > max_bytes:
        raise ValidationError(f"File exceeds maximum allowed size ({settings.MAX_UPLOAD_SIZE_MB} MB)")
        
    safe_name = sanitize_filename(filename)
    
    # Convert to PDF
    pdf_content = convert_to_pdf(safe_name, content)
    
    # Use sha256 of the *original* content for deduplication
    sha256_hash = compute_sha256(content)
    
    ws_raw_dir = settings.DATA_DIR / "workspaces" / workspace_id / "raw"
    ws_raw_dir.mkdir(parents=True, exist_ok=True)
    
    # Save the resulting PDF
    pdf_safe_name = safe_name if safe_name.lower().endswith(".pdf") else f"{safe_name}.pdf"
    target_path = ws_raw_dir / f"{sha256_hash[:12]}_{pdf_safe_name}"
    
    with open(target_path, "wb") as f:
        f.write(pdf_content)
        
    return target_path, pdf_safe_name, sha256_hash, len(pdf_content)
